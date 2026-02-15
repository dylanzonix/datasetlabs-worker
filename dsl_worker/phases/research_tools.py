"""
Research Tools - ChatGPT-style browsing for research agent.

Each ResearchTools instance has its OWN browser (one per scope).
Browser is lazy-initialized on first use, cleaned up when scope ends.

Tools:
- brave_search(query, response_length) → search results artifact
- open(ref_id_or_url, start_line, response_length) → page viewport
- find(ref_id, pattern, response_length) → matching lines  
- click(ref_id, link_id, response_length) → new page viewport
- note(content) → add to notes
- list_files(directory) → show available files with metadata
- code_exec(script) → execute Python with submit_seed()
- conclude_research(summary) → transition to decision mode
- breakdown(children) → split scope
- submit_seed(ref_id, lines, content) → create seed from source
- done(reason) → finish when seeds exhausted
- interact(url_or_ref_id, task) → browser agent for complex interactions
"""

import asyncio
import csv
import json
import logging
import random
import string
import time
from dataclasses import dataclass, field
from datetime import datetime, timezone
from enum import Enum
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional, Tuple

import httpx
from markdownify import markdownify as md

from dsl_worker.phases.artifacts import (
    ArtifactStore,
    SearchResults,
    SearchResult,
    PageView,
    PageLink,
    RESPONSE_LENGTHS,
    extract_links_from_markdown,
    format_viewport,
    format_links_table,
    format_search_results,
    find_in_lines,
)

logger = logging.getLogger(__name__)

# Browser-use imports
try:
    from browser_use import BrowserSession, Agent
    from browser_use.llm.openai.chat import ChatOpenAI
    BROWSER_AVAILABLE = True
except ImportError:
    BrowserSession = None
    Agent = None
    ChatOpenAI = None
    BROWSER_AVAILABLE = False


def short_id(length: int = 6) -> str:
    """Generate short alphanumeric ID."""
    return ''.join(random.choices(string.ascii_lowercase + string.digits, k=length))


class ResearchState(Enum):
    """State machine for research process."""
    RESEARCHING = "researching"
    CONCLUDED = "concluded"


@dataclass
class Seed:
    """A seed for row generation."""
    content: str
    scope_id: str
    scope_description: str
    notes: List[str]
    research_summary: Optional[str] = None
    source_ref: Optional[str] = None
    source_url: Optional[str] = None
    bucket_id: Optional[str] = None


@dataclass
class ResearchScope:
    """Current research scope context."""
    id: str
    description: str
    quota: int
    depth: int = 0
    notes: List[str] = field(default_factory=list)
    parent_notes: List[str] = field(default_factory=list)


# =============================================================================
# dsl_tools library injected into sandbox
# =============================================================================

DSL_TOOLS_LIBRARY = '''
"""
DSL Tools - Available in code execution sandbox.

Usage:
    from dsl_tools import submit_seed, list_files, file_info
"""

import json
import os
from pathlib import Path

_WORKSPACE = os.environ.get("DSL_WORKSPACE", "/workspace")
_SEEDS_FILE = os.path.join(_WORKSPACE, ".dsl_seeds.jsonl")


def submit_seed(content: str, source: str = None) -> None:
    """
    Submit a seed for row generation.
    
    Args:
        content: The seed content - what this row should be about
        source: Optional source reference (e.g., "data.pdf page 3")
    
    Example:
        submit_seed("Tesla Model 3", source="ev_list.csv row 15")
    """
    seed = {"content": content, "source": source}
    
    with open(_SEEDS_FILE, "a") as f:
        f.write(json.dumps(seed) + "\\n")
    
    print(f"[seed] {content[:60]}...")


def list_files(directory: str = "all") -> list:
    """List available files. Returns list of paths."""
    files = []
    
    if directory in ("uploads", "all"):
        uploads = Path(_WORKSPACE) / "uploads"
        if uploads.exists():
            files.extend([str(f) for f in uploads.iterdir() if f.is_file()])
    
    if directory in ("downloads", "all"):
        downloads = Path(_WORKSPACE) / "downloads"
        if downloads.exists():
            files.extend([str(f) for f in downloads.iterdir() if f.is_file()])
    
    return files


def file_info(path: str) -> dict:
    """Get file info: name, size_bytes, extension."""
    p = Path(path)
    if not p.exists():
        return {"exists": False}
    return {
        "exists": True,
        "name": p.name,
        "size_bytes": p.stat().st_size,
        "extension": p.suffix.lower(),
    }
'''


class ResearchTools:
    """
    Tools for research agent to explore and understand a domain.
    
    Each instance has its OWN browser (one per scope).
    
    State machine:
    - RESEARCHING: Can search, browse, take notes
    - CONCLUDED: Can breakdown or create seeds
    
    Must call conclude_research() before breakdown/seeding.
    """
    
    def __init__(
        self,
        workspace_dir: Path,
        schema: List[Dict],
        brave_api_key: Optional[str] = None,
        openai_client: Optional[Any] = None,
        model: str = "gpt-4o",
        sandbox: Optional[Any] = None,
        stop_checker: Optional[Callable[[], bool]] = None,
        blob_service_client: Optional[Any] = None,
        project_id: Optional[str] = None,
    ):
        self.workspace_dir = Path(workspace_dir)
        self.schema = schema
        self.brave_api_key = brave_api_key
        self.openai_client = openai_client
        self.model = model
        self.sandbox = sandbox
        self.stop_checker = stop_checker
        self.blob_service_client = blob_service_client
        self.project_id = str(project_id) if project_id else None
        
        # Ensure directories exist
        (self.workspace_dir / "uploads").mkdir(parents=True, exist_ok=True)
        (self.workspace_dir / "downloads").mkdir(parents=True, exist_ok=True)
        
        # Browser - lazy initialized, owned by this instance
        self._browser: Optional[Any] = None
        self._browser_init_lock = asyncio.Lock()
        self._open_pages: Dict[str, Any] = {}  # ref_id -> page
        
        # Artifact storage
        self.artifacts = ArtifactStore()
        
        # State machine
        self.state = ResearchState.RESEARCHING
        self.research_summary: Optional[str] = None
        self.research_actions: int = 0  # Track searches/opens/reads
        
        # Current scope
        self.scope: Optional[ResearchScope] = None
        
        # Track breakdown
        self.breakdown_children: Optional[List[Dict]] = None
        
        # Track seeds (in-memory)
        self.seeds: List[Seed] = []
        self.seeds_submitted: int = 0
        self.is_done: bool = False
    
    def set_scope(self, scope: ResearchScope):
        """Set current scope being researched."""
        self.scope = scope
        self.breakdown_children = None
        self.state = ResearchState.RESEARCHING
        self.research_summary = None
        self.research_actions = 0
        self.seeds = []
        self.seeds_submitted = 0
        self.is_done = False
    
    def _should_stop(self) -> bool:
        return self.stop_checker and self.stop_checker()
    
    def _get_config(self, response_length: str) -> Dict[str, int]:
        return RESPONSE_LENGTHS.get(response_length, RESPONSE_LENGTHS["medium"])
    
    def _require_state(self, required: ResearchState, action: str) -> Optional[str]:
        """Check state requirement, return error message if not met."""
        if self.state != required:
            if required == ResearchState.CONCLUDED:
                return f"Cannot {action} yet. You must call conclude_research() first to summarize what you learned."
            else:
                return f"Cannot {action} in {self.state.value} state."
        return None
    
    # =========================================================================
    # Browser Management
    # =========================================================================
    
    async def _get_browser(self) -> Any:
        """Get or create BrowserSession for this scope.

        Uses a lock so parallel open() calls don't race through init
        (the browser is assigned before start() awaits, so without a lock
        a second caller would see a non-None but not-yet-started session).
        """
        if not BROWSER_AVAILABLE:
            raise RuntimeError("browser-use not installed")

        async with self._browser_init_lock:
            if self._browser is None:
                from dsl_worker.config import settings

                # Build proxy settings if configured (Bright Data residential)
                proxy = None
                if settings.browser_proxy_server:
                    from browser_use.browser import ProxySettings
                    proxy = ProxySettings(
                        server=settings.browser_proxy_server,
                        username=settings.browser_proxy_username or None,
                        password=settings.browser_proxy_password or None,
                    )

                # Load cookies from Azure Blob (global + per-project)
                storage_state = None
                if self.blob_service_client and self.project_id:
                    from dsl_worker.phases.cookie_manager import load_cookies
                    storage_state = load_cookies(
                        self.blob_service_client,
                        settings.azure_storage_container_name,
                        self.project_id,
                        settings.browser_global_cookies_blob_path,
                    )

                self._browser = BrowserSession(
                    headless=False,
                    downloads_path=str(self.workspace_dir / "downloads"),
                    auto_download_pdfs=True,
                    keep_alive=True,
                    proxy=proxy,
                    storage_state=storage_state,
                )
                await self._browser.start()
                scope_id = self.scope.id if self.scope else "unknown"
                logger.info(f"[ResearchTools] BrowserSession started for scope {scope_id}")

        return self._browser
    
    async def cleanup(self):
        """Cleanup browser session. Saves cookies, then stops browser."""
        if self._browser:
            try:
                # Save project cookies to Azure Blob before stopping
                if self.blob_service_client and self.project_id:
                    from dsl_worker.config import settings
                    from dsl_worker.phases.cookie_manager import save_project_cookies
                    try:
                        save_project_cookies(
                            self._browser,
                            self.blob_service_client,
                            settings.azure_storage_container_name,
                            self.project_id,
                        )
                    except Exception as e:
                        logger.warning(f"[ResearchTools] Failed to save cookies: {e}")

                await self._browser.stop()
                scope_id = self.scope.id if self.scope else "unknown"
                logger.info(f"[ResearchTools] BrowserSession closed for scope {scope_id}")
            except Exception as e:
                logger.warning(f"[ResearchTools] Error stopping BrowserSession: {e}")
            self._browser = None
    
    # =========================================================================
    # brave_search
    # =========================================================================
    
    async def brave_search(self, query: str, response_length: str = "medium") -> Tuple[str, float]:
        """Search the web using Brave Search API with retry."""
        if not self.brave_api_key:
            return "Error: Brave API key not configured", 0.0
        
        config = self._get_config(response_length)
        count = config["results"]
        
        max_retries = 3
        base_delay = 1.0
        
        for attempt in range(max_retries):
            try:
                async with httpx.AsyncClient() as client:
                    resp = await client.get(
                        "https://api.search.brave.com/res/v1/web/search",
                        params={"q": query, "count": count},
                        headers={"X-Subscription-Token": self.brave_api_key},
                        timeout=30.0,
                    )
                    resp.raise_for_status()
                    data = resp.json()
                
                results = []
                for i, r in enumerate(data.get("web", {}).get("results", [])):
                    results.append(SearchResult(
                        id=i,
                        title=r.get("title", ""),
                        url=r.get("url", ""),
                        snippet=r.get("description", ""),
                        date=r.get("age"),
                    ))
                
                artifact = SearchResults(query=query, results=results)
                ref_id = self.artifacts.store_search(artifact)
                
                return format_search_results(artifact, ref_id, count), 0.0
                
            except httpx.HTTPStatusError as e:
                if e.response.status_code == 429 or e.response.status_code >= 500:
                    if attempt < max_retries - 1:
                        delay = base_delay * (2 ** attempt)
                        logger.warning(
                            f"[ResearchTools] Brave search failed ({e.response.status_code}), "
                            f"retry {attempt + 1}/{max_retries} in {delay}s"
                        )
                        await asyncio.sleep(delay)
                        continue
                logger.error(f"[ResearchTools] Brave search HTTP error: {e}")
                return f"Search error: {e}", 0.0
                
            except Exception as e:
                if attempt < max_retries - 1:
                    delay = base_delay * (2 ** attempt)
                    logger.warning(
                        f"[ResearchTools] Brave search failed ({e}), "
                        f"retry {attempt + 1}/{max_retries} in {delay}s"
                    )
                    await asyncio.sleep(delay)
                    continue
                logger.error(f"[ResearchTools] brave_search failed: {e}")
                return f"Search error: {e}", 0.0
        
        return "Search failed after retries", 0.0
    
    # =========================================================================
    # open
    # =========================================================================
    
    # Max time for a single open() call (navigation + content extraction)
    OPEN_TIMEOUT = 60

    async def open(
        self,
        ref_id_or_url: str,
        start_line: int = 0,
        response_length: str = "medium",
    ) -> Tuple[str, float]:
        """Open a URL or navigate within existing page artifact.

        Each call opens a NEW browser tab so multiple open() calls can run
        in parallel without fighting over the same page. Tabs are closed
        when done. All tabs share the same BrowserSession (shared cookies).
        """
        config = self._get_config(response_length)
        num_lines = config["lines"]

        artifact = self.artifacts.get(ref_id_or_url)

        if artifact:
            if isinstance(artifact, PageView):
                viewport = format_viewport(
                    artifact.lines, start_line, num_lines,
                    ref_id_or_url, artifact.url
                )
                links = format_links_table(artifact.links)
                return f"{viewport}\n{links}", 0.0

            elif isinstance(artifact, SearchResults):
                return format_search_results(artifact, ref_id_or_url, config["results"]), 0.0

            else:
                return f"Unknown artifact type for: {ref_id_or_url}", 0.0

        url = ref_id_or_url
        if not url.startswith(("http://", "https://")):
            url = "https://" + url

        try:
            return await asyncio.wait_for(
                self._open_in_new_tab(url, start_line, num_lines),
                timeout=self.OPEN_TIMEOUT,
            )
        except asyncio.TimeoutError:
            logger.warning(f"[ResearchTools] open() timed out after {self.OPEN_TIMEOUT}s for {url}")
            return f"Timed out loading {url} after {self.OPEN_TIMEOUT}s", 0.0
        except Exception as e:
            logger.error(f"[ResearchTools] open failed for {url}: {e}")
            return f"Failed to open {url}: {e}", 0.0

    async def _open_in_new_tab(
        self, url: str, start_line: int, num_lines: int,
    ) -> Tuple[str, float]:
        """Open a URL in a new browser tab, extract content, close tab."""
        session = await self._get_browser()
        page = None

        try:
            # Track downloads before navigation (session-level)
            downloads_before = set(session.downloaded_files) if session.downloaded_files else set()

            # Create a blank tab, then navigate with proper load-waiting.
            # new_page(url) fires-and-forgets; page.goto() waits for
            # networkIdle/load lifecycle events so fast pages return fast
            # and slow pages get the time they need.
            page = await session.new_page()
            try:
                await page.goto(url)
            except Exception as nav_error:
                # Timeout is expected for PDF downloads — continue and check
                logger.debug(f"[ResearchTools] Navigation exception (may be expected): {nav_error}")

            # Check for new downloads
            downloads_after = set(session.downloaded_files) if session.downloaded_files else set()
            new_downloads = downloads_after - downloads_before

            if new_downloads:
                return await self._handle_download(
                    new_downloads, url, start_line, num_lines,
                )

            # Try to extract HTML content from this tab
            markdown = ""
            max_attempts = 4

            for attempt in range(max_attempts):
                try:
                    html = await page.evaluate('() => document.body.innerHTML')
                    if html and len(html.strip()) > 0:
                        markdown = md(html, heading_style='ATX')
                        break
                except Exception as e:
                    logger.debug(f"[ResearchTools] HTML extraction attempt {attempt + 1} failed: {e}")
                await asyncio.sleep(1.5)

            if not markdown:
                # One more check for late downloads (e.g. PDF redirect)
                await asyncio.sleep(2.0)
                downloads_after = set(session.downloaded_files) if session.downloaded_files else set()
                new_downloads = downloads_after - downloads_before

                if new_downloads:
                    return await self._handle_download(
                        new_downloads, url, start_line, num_lines,
                    )

                markdown = "Page loaded but no content extracted"

            lines = markdown.split('\n')
            links = extract_links_from_markdown(markdown, url)

            page_view = PageView(
                url=url,
                fetched_at=datetime.now(timezone.utc).isoformat(),
                lines=lines,
                total_lines=len(lines),
                links=links,
            )
            ref_id = self.artifacts.store_page(page_view)

            viewport = format_viewport(lines, start_line, num_lines, ref_id, url)
            links_table = format_links_table(links)

            return f"{viewport}\n{links_table}", 0.0

        finally:
            # Always close the tab to prevent leaks
            if page is not None:
                try:
                    await session.close_page(page)
                except Exception:
                    pass

    def _upload_download_to_blob(self, local_path: Path) -> Optional[str]:
        """Upload a downloaded file to Azure Blob for persistence. Returns blob path or None."""
        if not self.blob_service_client or not self.project_id:
            return None

        from dsl_worker.config import settings

        blob_path = f"projects/{self.project_id}/downloads/{local_path.name}"
        try:
            blob_client = self.blob_service_client.get_blob_client(
                container=settings.azure_storage_container_name,
                blob=blob_path,
            )
            with open(local_path, "rb") as f:
                blob_client.upload_blob(f, overwrite=True)
            logger.info(f"[ResearchTools] Uploaded download to blob: {blob_path}")
            return blob_path
        except Exception as e:
            logger.warning(f"[ResearchTools] Failed to upload download to blob: {e}")
            return None

    async def _handle_download(
        self,
        new_downloads: set,
        url: str,
        start_line: int,
        num_lines: int,
    ) -> Tuple[str, float]:
        """Process a downloaded file detected during open()."""
        downloaded_path = Path(list(new_downloads)[0])
        logger.info(f"[ResearchTools] File downloaded: {downloaded_path.name}")

        # Wait a bit for large files to finish writing
        await asyncio.sleep(1.0)

        # Upload to Azure Blob for durability (file also stays local for immediate use)
        self._upload_download_to_blob(downloaded_path)

        content, file_info = await self._extract_file_content(downloaded_path)

        if content:
            lines = content.split('\n')

            page_view = PageView(
                url=url,
                fetched_at=datetime.now(timezone.utc).isoformat(),
                lines=lines,
                total_lines=len(lines),
                links=[],
            )
            ref_id = self.artifacts.store_page(page_view)

            header = f"Downloaded: {downloaded_path.name} ({file_info})\n"
            header += f"Source: {url}\n"
            header += "=" * 50 + "\n"

            viewport = format_viewport(lines, start_line, num_lines, ref_id, url)
            return f"{header}{viewport}", 0.0
        else:
            size_kb = downloaded_path.stat().st_size / 1024
            return (
                f"File downloaded: {downloaded_path.name} ({size_kb:.1f}KB)\n"
                f"Location: /workspace/downloads/{downloaded_path.name}\n"
                f"Source: {url}\n\n"
                f"Could not extract text. Use code_exec() with appropriate library to process."
            ), 0.0
    
    async def _extract_file_content(self, file_path: Path) -> Tuple[str, str]:
        """
        Extract text content from a downloaded file.
        
        Returns (content, file_info) tuple.
        content is the extracted text, or empty string if extraction failed.
        file_info is a brief description like "PDF, 15 pages" or "245KB".
        """
        ext = file_path.suffix.lower()
        size_kb = file_path.stat().st_size / 1024
        
        try:
            if ext == '.pdf':
                try:
                    import pdfplumber
                    text_lines = []
                    page_count = 0
                    
                    with pdfplumber.open(file_path) as pdf:
                        page_count = len(pdf.pages)
                        for i, page in enumerate(pdf.pages):
                            text = page.extract_text() or ""
                            if text.strip():
                                text_lines.append(f"--- Page {i + 1} ---")
                                text_lines.extend(text.split('\n'))
                    
                    if text_lines:
                        return '\n'.join(text_lines), f"PDF, {page_count} pages"
                    else:
                        return "", f"PDF, {page_count} pages (no extractable text)"
                        
                except ImportError:
                    logger.warning("[ResearchTools] pdfplumber not installed")
                    return "", f"PDF, {size_kb:.1f}KB (pdfplumber not installed)"
            
            elif ext in ('.txt', '.md', '.csv', '.json', '.xml', '.html', '.htm'):
                content = file_path.read_text(errors='ignore')
                line_count = len(content.split('\n'))
                return content, f"{ext[1:].upper()}, {line_count} lines"
            
            elif ext in ('.xlsx', '.xls'):
                try:
                    import pandas as pd
                    
                    # Read all sheets
                    xlsx = pd.ExcelFile(file_path)
                    all_text = []
                    
                    for sheet_name in xlsx.sheet_names:
                        df = pd.read_excel(xlsx, sheet_name=sheet_name)
                        all_text.append(f"--- Sheet: {sheet_name} ---")
                        all_text.append(df.to_string())
                    
                    return '\n'.join(all_text), f"Excel, {len(xlsx.sheet_names)} sheets"
                    
                except ImportError:
                    return "", f"Excel, {size_kb:.1f}KB (pandas not installed)"
            
            elif ext == '.docx':
                try:
                    from docx import Document
                    doc = Document(file_path)
                    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
                    return '\n\n'.join(paragraphs), f"Word doc, {len(paragraphs)} paragraphs"
                except ImportError:
                    return "", f"Word doc, {size_kb:.1f}KB (python-docx not installed)"
            
            else:
                # Unknown type - can't extract
                return "", f"{size_kb:.1f}KB"
                
        except Exception as e:
            logger.error(f"[ResearchTools] File extraction failed for {file_path}: {e}")
            return "", f"{size_kb:.1f}KB (extraction error: {e})"
    
    # =========================================================================
    # find
    # =========================================================================
    
    async def find(
        self,
        ref_id: str,
        pattern: str,
        response_length: str = "medium",
    ) -> Tuple[str, float]:
        """Find pattern in a page artifact."""
        page = self.artifacts.get_page(ref_id)
        if not page:
            return f"Page not found: {ref_id}", 0.0
        
        config = self._get_config(response_length)
        return find_in_lines(page.lines, pattern, config["matches"]), 0.0
    
    # =========================================================================
    # click
    # =========================================================================
    
    async def click(
        self,
        ref_id: str,
        link_id: int,
        response_length: str = "medium",
    ) -> Tuple[str, float]:
        """Click a link in a page artifact."""
        page = self.artifacts.get_page(ref_id)
        if not page:
            return f"Page not found: {ref_id}", 0.0
        
        link = None
        for l in page.links:
            if l.id == link_id:
                link = l
                break
        
        if not link:
            return f"Link {link_id} not found in {ref_id}. Available: 0-{len(page.links)-1}", 0.0
        
        return await self.open(link.href, start_line=0, response_length=response_length)
    
    # =========================================================================
    # note
    # =========================================================================
    
    def note(self, content: str) -> Tuple[str, float]:
        """Record a note about this domain."""
        if not self.scope:
            return "No active scope", 0.0
        
        self.scope.notes.append(content)
        return f"Noted ({len(self.scope.notes)} total)", 0.0
    
    # =========================================================================
    # list_files
    # =========================================================================
    
    async def list_files(self, directory: str = "all") -> Tuple[str, float]:
        """List available files with metadata preview."""
        output = []
        
        for subdir in ["uploads", "downloads"]:
            if directory not in (subdir, "all"):
                continue
            
            dir_path = self.workspace_dir / subdir
            if not dir_path.exists():
                continue
            
            files = sorted([f for f in dir_path.iterdir() if f.is_file()])
            if not files:
                continue
            
            icon = "📁" if subdir == "uploads" else "📥"
            output.append(f"{icon} {subdir.title()}:")
            
            for f in files:
                size_kb = f.stat().st_size / 1024
                meta = self._get_file_metadata(f)
                
                output.append(f"  {f.name} ({size_kb:.1f}KB)")
                if meta:
                    output.append(f"    └─ {meta}")
        
        if not output:
            return "No files found. Upload files or browse web to download.", 0.0
        
        output.append("\nUse code_exec() to parse files and submit_seed() for items.")
        return '\n'.join(output), 0.0
    
    def _get_file_metadata(self, path: Path) -> str:
        """Quick metadata preview for common file types."""
        ext = path.suffix.lower()
        
        try:
            if ext == '.csv':
                with open(path, 'r', errors='ignore') as f:
                    reader = csv.reader(f)
                    header = next(reader, None)
                    row_count = sum(1 for _ in f) + 1
                if header:
                    cols = header[:5]
                    more = f" +{len(header)-5} more" if len(header) > 5 else ""
                    return f"Columns: {cols}{more} | ~{row_count} rows"
                return f"~{row_count} rows"
            
            elif ext == '.json':
                text = path.read_text(errors='ignore')[:10000]
                data = json.loads(text)
                if isinstance(data, list):
                    return f"Array with {len(data)} items"
                elif isinstance(data, dict):
                    keys = list(data.keys())[:5]
                    more = f" +{len(data)-5}" if len(data) > 5 else ""
                    return f"Object keys: {keys}{more}"
            
            elif ext in ('.xlsx', '.xls'):
                return "Excel spreadsheet"
            
            elif ext == '.pdf':
                return "PDF document"
            
            elif ext in ('.db', '.sqlite', '.sqlite3'):
                return "SQLite database"
            
            elif ext in ('.txt', '.md'):
                lines = len(path.read_text(errors='ignore').split('\n'))
                return f"~{lines} lines"
                
        except Exception:
            pass
        
        return ""
    
    # =========================================================================
    # code_exec
    # =========================================================================
    
    async def code_exec(self, script: str, description: str = "") -> Tuple[str, float]:
        """
        Execute Python code with access to files and submit_seed().
        
        Available in script:
        - submit_seed(content, source=None) - Submit a seed for row generation
        - list_files() - List available files
        - file_info(path) - Get file metadata
        
        Files are at:
        - /workspace/uploads/ - User uploaded files
        - /workspace/downloads/ - Browser downloaded files
        """
        if not self.sandbox:
            return "Code execution not available", 0.0
        
        # Clear previous seeds file
        seeds_file = self.workspace_dir / ".dsl_seeds.jsonl"
        if seeds_file.exists():
            seeds_file.unlink()
        
        # Build code with dsl_tools injected
        full_script = f"""
import os
os.environ["DSL_WORKSPACE"] = "{self.workspace_dir}"

{DSL_TOOLS_LIBRARY}

# Common libraries
try:
    import pandas as pd
except ImportError:
    pd = None

try:
    import pdfplumber
except ImportError:
    pdfplumber = None

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

# User script
{script}
"""
        
        result = self.sandbox.execute(
            script=full_script,
            workspace_dir=str(self.workspace_dir),
            timeout=120,
        )
        
        # Read any seeds that were submitted
        new_seeds = []
        if seeds_file.exists():
            for line in seeds_file.read_text().strip().split('\n'):
                if line:
                    try:
                        seed_data = json.loads(line)
                        new_seeds.append(seed_data)
                    except json.JSONDecodeError:
                        pass
        
        # Add seeds to our collection
        for seed_data in new_seeds:
            self._add_seed_from_code(seed_data)
        
        # Build response
        output_parts = []
        
        if result.stdout:
            output_parts.append(result.stdout)
        
        if not result.success and result.stderr:
            output_parts.append(f"Error: {result.stderr}")
        
        if new_seeds:
            output_parts.append(f"\n✓ {len(new_seeds)} seeds submitted")
            remaining = self.scope.quota - self.seeds_submitted if self.scope else 0
            output_parts.append(f"  Remaining quota: {remaining}")
        
        return '\n'.join(output_parts) if output_parts else "Code executed (no output)", 0.0
    
    def _add_seed_from_code(self, seed_data: Dict):
        """Add a seed that was submitted via code execution."""
        if not self.scope:
            return
        
        seed = Seed(
            content=seed_data.get("content", ""),
            scope_id=self.scope.id,
            scope_description=self.scope.description,
            notes=self.scope.parent_notes + self.scope.notes,
            research_summary=self.research_summary,
            source_ref=seed_data.get("source"),
            source_url=None,
        )
        self.seeds.append(seed)
        self.seeds_submitted += 1
    
    # =========================================================================
    # conclude_research
    # =========================================================================
    
    def conclude_research(self, summary: str) -> Tuple[str, float]:
        """
        Conclude research phase and transition to decision mode.
        
        Must be called before breakdown() or submit_seed().
        Requires actual research to have been performed.
        """
        if not self.scope:
            return "No active scope", 0.0
        
        if not summary or len(summary.strip()) < 10:
            return "Please provide a meaningful summary of your research findings.", 0.0
        
        if self.research_actions == 0:
            return (
                "Cannot conclude research without having done any. "
                "You need to use your research tools (brave_search, open, etc.) "
                "to build understanding before concluding. Start by searching broadly "
                "to understand your scope's domain."
            ), 0.0
        
        self.state = ResearchState.CONCLUDED
        self.research_summary = summary
        
        return (
            f"Research concluded. Summary recorded.\n\n"
            f"You can now either:\n"
            f"- breakdown(children) to split into smaller scopes\n"
            f"- submit_seed() to submit seeds for row generation\n"
            f"- done(reason) if seeds are exhausted before quota"
        ), 0.0
    
    # =========================================================================
    # breakdown
    # =========================================================================
    
    def breakdown(self, children: List[Dict]) -> Tuple[str, float]:
        """Break scope into smaller sub-scopes."""
        error = self._require_state(ResearchState.CONCLUDED, "breakdown")
        if error:
            return error, 0.0
        
        if not self.scope:
            return "No active scope", 0.0
        
        if not children:
            return "No children provided", 0.0
        
        self.breakdown_children = children
        
        total_weight = sum(c.get("weight", 1.0) for c in children)
        lines = [f"Breaking into {len(children)} sub-scopes:"]
        
        for c in children:
            weight = c.get("weight", 1.0) / total_weight
            approx_quota = int(self.scope.quota * weight)
            lines.append(f"  - {c.get('description', '?')} (~{approx_quota} rows)")
        
        return '\n'.join(lines), 0.0
    
    # =========================================================================
    # submit_seed
    # =========================================================================
    
    def submit_seed(
        self,
        ref_id: Optional[str] = None,
        lines: Optional[List[int]] = None,
        content: Optional[str] = None,
    ) -> Tuple[str, float]:
        """Submit a single seed for row generation."""
        error = self._require_state(ResearchState.CONCLUDED, "submit seed")
        if error:
            return error, 0.0
        
        if not self.scope:
            return "No active scope", 0.0
        
        # Build seed content
        seed_content = ""
        source_ref = None
        source_url = None
        
        # Extract from source if ref_id + lines provided
        if ref_id and lines and len(lines) == 2:
            page = self.artifacts.get_page(ref_id)
            if not page:
                return f"Page not found: {ref_id}", 0.0
            
            start, end = lines
            start = max(0, start)
            end = min(len(page.lines), end + 1)
            extracted = '\n'.join(page.lines[start:end]).strip()
            
            if extracted:
                seed_content = extracted
                source_ref = ref_id
                source_url = page.url
        
        # Add written content
        if content:
            if seed_content:
                seed_content = f"{seed_content}\n\n{content}"
            else:
                seed_content = content
        
        if not seed_content.strip():
            return "Empty seed - provide ref_id+lines and/or content", 0.0
        
        # Create seed
        all_notes = self.scope.parent_notes + self.scope.notes
        seed = Seed(
            content=seed_content,
            scope_id=self.scope.id,
            scope_description=self.scope.description,
            notes=all_notes,
            research_summary=self.research_summary,
            source_ref=source_ref,
            source_url=source_url,
        )
        self.seeds.append(seed)
        self.seeds_submitted += 1
        
        remaining = self.scope.quota - self.seeds_submitted
        
        if remaining > 0:
            return f"Seed {self.seeds_submitted} submitted | Remaining quota: {remaining}", 0.0
        else:
            return f"Seed {self.seeds_submitted} submitted | Quota filled!", 0.0
    
    # =========================================================================
    # done
    # =========================================================================
    
    def done(self, reason: str) -> Tuple[str, float]:
        """Finish when seeds are exhausted before reaching quota."""
        error = self._require_state(ResearchState.CONCLUDED, "done")
        if error:
            return error, 0.0
        
        if not self.scope:
            return "No active scope", 0.0
        
        if not reason or len(reason.strip()) < 5:
            return "Please provide a reason why seeds are exhausted", 0.0
        
        self.is_done = True
        remaining = self.scope.quota - self.seeds_submitted
        
        return f"Done. Submitted {self.seeds_submitted} seeds. Remaining {remaining} could not be filled. Reason: {reason}", 0.0
    
    @property
    def remaining_quota(self) -> int:
        """Get remaining quota for this scope."""
        if not self.scope:
            return 0
        return max(0, self.scope.quota - self.seeds_submitted)
    
    @property
    def quota_filled(self) -> bool:
        """Check if quota is filled."""
        return self.remaining_quota == 0
    
    # =========================================================================
    # interact (Browser Agent)
    # =========================================================================
    
    async def interact(self, url_or_ref_id: str, task: str) -> Tuple[str, float]:
        """
        Use Browser Agent for complex interactions on a page.

        Uses the SAME browser session as open() — session persists after
        completion (keep_alive=True), so cookies, auth state, and page
        context are preserved for subsequent open()/click() calls.
        """
        if not BROWSER_AVAILABLE:
            return "browser-use not installed", 0.0

        if not self.openai_client:
            return "OpenAI client not initialized for interact()", 0.0

        # Get our browser session
        session = await self._get_browser()

        # Resolve URL
        url = url_or_ref_id
        page_artifact = self.artifacts.get_page(url_or_ref_id)
        if page_artifact:
            url = page_artifact.url
        elif not url.startswith(("http://", "https://")):
            url = "https://" + url

        total_cost = 0.0

        try:
            # Create browser-use LLM
            browser_llm = ChatOpenAI(model=self.model)

            agent = Agent(
                task=f"Navigate to {url} and: {task}",
                browser_session=session,
                llm=browser_llm,
                calculate_cost=True,
            )

            history = await agent.run(max_steps=30)

            # Get cost
            if hasattr(history, 'usage') and history.usage:
                if hasattr(history.usage, 'total_cost'):
                    total_cost += history.usage.total_cost

            # Clean up extra tabs the agent may have opened (prevent leaks)
            try:
                pages = await session.get_pages()
                if len(pages) > 1:
                    for extra_page in pages[1:]:
                        try:
                            await session.close_page(extra_page)
                        except Exception:
                            pass
            except Exception:
                pass

            # Capture final page state as artifact
            try:
                current_page = await session.get_current_page()
                if current_page:
                    current_url = await current_page.evaluate('() => window.location.href')
                    html = await current_page.evaluate('() => document.body.innerHTML')
                    markdown = md(html, heading_style='ATX')

                    lines = markdown.split('\n')
                    links = extract_links_from_markdown(markdown, current_url)

                    page_view = PageView(
                        url=current_url,
                        fetched_at=datetime.now(timezone.utc).isoformat(),
                        lines=lines,
                        total_lines=len(lines),
                        links=links,
                    )
                    final_ref_id = self.artifacts.store_page(page_view)

                    return f"Browser session completed. Final page stored as {final_ref_id}", total_cost
            except Exception as e:
                logger.warning(f"[interact] Could not capture final state: {e}")

            return "Browser session completed", total_cost

        except Exception as e:
            logger.error(f"[interact] Browser agent failed: {e}")
            return f"Browser agent error: {e}", total_cost
    
    # =========================================================================
    # Tool Registration for Agent Framework
    # =========================================================================

    def register_on(self, registry: 'ToolRegistry') -> None:
        """
        Register browsing/research tools onto an agent ToolRegistry.

        Registers: brave_search, open, find, click, list_files, code_exec,
        interact. Skips scope-dependent tools (note, conclude_research,
        submit_seed, breakdown, done).

        This is the single place adapter functions live — agents call this
        instead of duplicating the boilerplate.
        """
        # Adapter closures that translate args dict -> method calls
        handlers = {
            "brave_search": lambda args: self.brave_search(
                query=args.get("query", ""),
                response_length=args.get("response_length", "medium"),
            ),
            "open": lambda args: self.open(
                ref_id_or_url=args.get("ref_id_or_url", ""),
                start_line=args.get("start_line", 0),
                response_length=args.get("response_length", "medium"),
            ),
            "find": lambda args: self.find(
                ref_id=args.get("ref_id", ""),
                pattern=args.get("pattern", ""),
                response_length=args.get("response_length", "medium"),
            ),
            "click": lambda args: self.click(
                ref_id=args.get("ref_id", ""),
                link_id=args.get("link_id", 0),
                response_length=args.get("response_length", "medium"),
            ),
            "list_files": lambda args: self.list_files(
                directory=args.get("directory", "all"),
            ),
            "code_exec": lambda args: self.code_exec(
                script=args.get("script", ""),
                description=args.get("description", ""),
            ),
            "interact": lambda args: self.interact(
                url_or_ref_id=args.get("url_or_ref_id", ""),
                task=args.get("task", ""),
            ),
        }

        for defn in self._research_tools():
            name = defn.get("name")
            if name in handlers:
                registry.add(
                    name=name,
                    description=defn.get("description", ""),
                    parameters=defn.get("parameters", {}),
                    handler=handlers[name],
                )

    # =========================================================================
    # Tool Definitions for LLM
    # =========================================================================

    def get_tool_definitions(self, phase: str = None) -> List[Dict]:
        """Get tool definitions for current or specified phase."""
        if phase is None:
            phase = "research" if self.state == ResearchState.RESEARCHING else "decision"
        
        if phase == "research":
            return self._research_tools()
        else:
            return self._decision_tools()
    
    def _research_tools(self) -> List[Dict]:
        """Tools available during research phase."""
        return [
            {
                "type": "function",
                "name": "brave_search",
                "description": "Search the web. Returns results with URLs you can open.",
                "parameters": {
                    "type": "object",
                    "properties": {
                        "query": {
                            "type": "string",
                            "description": "Search query"
                        },
                        "response_length": {
                            "type": "string",
                            "enum": ["short", "medium", "long"],
                            "description": "short=5, medium=10, long=20 results"
                        },
                    },
                    "required": ["query"]
                }
            },
            {
                "type": "function",
                "name": "open",
                "description": "Open a URL or view lines from existing page. Returns line-numbered content.",
                "parameters": {
                    "type": "object",
                    "properties": {
                        "ref_id_or_url": {
                            "type": "string",
                            "description": "URL to fetch, or ref_id (p0, p1...) to view existing"
                        },
                        "start_line": {
                            "type": "integer",
                            "description": "Line to start from (default: 0)"
                        },
                        "response_length": {
                            "type": "string",
                            "enum": ["short", "medium", "long"],
                            "description": "short=60, medium=150, long=300 lines"
                        },
                    },
                    "required": ["ref_id_or_url"]
                }
            },
            {
                "type": "function",
                "name": "find",
                "description": "Find text/pattern in a page. Returns matching lines with numbers.",
                "parameters": {
                    "type": "object",
                    "properties": {
                        "ref_id": {
                            "type": "string",
                            "description": "Page ref_id (p0, p1...)"
                        },
                        "pattern": {
                            "type": "string",
                            "description": "Text or regex to find"
                        },
                        "response_length": {
                            "type": "string",
                            "enum": ["short", "medium", "long"],
                            "description": "Controls max matches returned"
                        },
                    },
                    "required": ["ref_id", "pattern"]
                }
            },
            {
                "type": "function",
                "name": "click",
                "description": "Click a link in a page. Opens the linked URL.",
                "parameters": {
                    "type": "object",
                    "properties": {
                        "ref_id": {
                            "type": "string",
                            "description": "Page ref_id"
                        },
                        "link_id": {
                            "type": "integer",
                            "description": "Link ID from links table"
                        },
                        "response_length": {
                            "type": "string",
                            "enum": ["short", "medium", "long"],
                            "description": "Viewport size for new page"
                        },
                    },
                    "required": ["ref_id", "link_id"]
                }
            },
            {
                "type": "function",
                "name": "note",
                "description": "Record a note about this domain. Use freely for facts, patterns, constraints, observations.",
                "parameters": {
                    "type": "object",
                    "properties": {
                        "content": {
                            "type": "string",
                            "description": "What you observed or learned"
                        },
                    },
                    "required": ["content"]
                }
            },
            {
                "type": "function",
                "name": "list_files",
                "description": "List available files (uploads and downloads) with metadata preview.",
                "parameters": {
                    "type": "object",
                    "properties": {
                        "directory": {
                            "type": "string",
                            "enum": ["uploads", "downloads", "all"],
                            "description": "Which directory to list (default: all)"
                        }
                    }
                }
            },
            {
                "type": "function",
                "name": "code_exec",
                "description": """Execute Python code with file access.

Available:
- list_files(), file_info(path) - File utilities
- pandas, pdfplumber, openpyxl (if installed)

Files at /workspace/uploads/ and /workspace/downloads/""",
                "parameters": {
                    "type": "object",
                    "properties": {
                        "script": {
                            "type": "string",
                            "description": "Python code to execute"
                        },
                        "description": {
                            "type": "string",
                            "description": "What this code does"
                        }
                    },
                    "required": ["script"]
                }
            },
            {
                "type": "function",
                "name": "conclude_research",
                "description": "Conclude research phase and transition to decision phase. Summarize what you learned about the domain. After this, your tools will change to breakdown/seeding tools.",
                "parameters": {
                    "type": "object",
                    "properties": {
                        "summary": {
                            "type": "string",
                            "description": "Summary of research findings and domain understanding"
                        },
                    },
                    "required": ["summary"]
                }
            },
            {
                "type": "function",
                "name": "interact",
                "description": "Use Browser Agent for complex page interactions (login, forms, captchas, JS-heavy pages, pagination). Browser session persists — cookies and state are preserved for subsequent open()/click() calls.",
                "parameters": {
                    "type": "object",
                    "properties": {
                        "url_or_ref_id": {
                            "type": "string",
                            "description": "URL or ref_id to interact with"
                        },
                        "task": {
                            "type": "string",
                            "description": "What to accomplish (e.g., 'login', 'click Load More')"
                        },
                    },
                    "required": ["url_or_ref_id", "task"]
                }
            },
        ]
    
    def _decision_tools(self) -> List[Dict]:
        """Tools available during decision/seeding phase."""
        return [
            {
                "type": "function",
                "name": "breakdown",
                "description": "Break scope into sub-scopes. Each child becomes its own research agent.",
                "parameters": {
                    "type": "object",
                    "properties": {
                        "children": {
                            "type": "array",
                            "items": {
                                "type": "object",
                                "properties": {
                                    "description": {"type": "string"},
                                    "weight": {"type": "number"},
                                },
                                "required": ["description"]
                            },
                            "description": "Sub-scopes with descriptions and relative weights"
                        },
                    },
                    "required": ["children"]
                }
            },
            {
                "type": "function",
                "name": "submit_seed",
                "description": "Submit a seed for row generation. Provide source extraction (ref_id + lines), written content, or both.",
                "parameters": {
                    "type": "object",
                    "properties": {
                        "ref_id": {
                            "type": "string",
                            "description": "Page ref_id to extract from (optional)"
                        },
                        "lines": {
                            "type": "array",
                            "items": {"type": "integer"},
                            "minItems": 2,
                            "maxItems": 2,
                            "description": "[start, end] line range to extract (requires ref_id)"
                        },
                        "content": {
                            "type": "string",
                            "description": "Written seed content or additional context (optional)"
                        },
                    }
                }
            },
            {
                "type": "function",
                "name": "done",
                "description": "Finish when seeds are exhausted before reaching quota. Use when seeds are finite (real items that exist or don't) and you've found all that exist. Don't use to quit early.",
                "parameters": {
                    "type": "object",
                    "properties": {
                        "reason": {
                            "type": "string",
                            "description": "Why seeds are exhausted"
                        },
                    },
                    "required": ["reason"]
                }
            },
            {
                "type": "function",
                "name": "note",
                "description": "Record a note. Use to track coverage while seeding.",
                "parameters": {
                    "type": "object",
                    "properties": {
                        "content": {
                            "type": "string",
                            "description": "What you observed or decided"
                        },
                    },
                    "required": ["content"]
                }
            },
            {
                "type": "function",
                "name": "brave_search",
                "description": "Search the web. Still available if you need to check something while seeding.",
                "parameters": {
                    "type": "object",
                    "properties": {
                        "query": {
                            "type": "string",
                            "description": "Search query"
                        },
                        "response_length": {
                            "type": "string",
                            "enum": ["short", "medium", "long"],
                            "description": "short=5, medium=10, long=20 results"
                        },
                    },
                    "required": ["query"]
                }
            },
            {
                "type": "function",
                "name": "open",
                "description": "Open a URL or view lines from existing page. Still available for reference.",
                "parameters": {
                    "type": "object",
                    "properties": {
                        "ref_id_or_url": {
                            "type": "string",
                            "description": "URL to fetch, or ref_id (p0, p1...) to view existing"
                        },
                        "start_line": {
                            "type": "integer",
                            "description": "Line to start from (default: 0)"
                        },
                        "response_length": {
                            "type": "string",
                            "enum": ["short", "medium", "long"],
                            "description": "short=60, medium=150, long=300 lines"
                        },
                    },
                    "required": ["ref_id_or_url"]
                }
            },
        ]
    
    async def execute_tool(self, name: str, args: Dict) -> Tuple[str, float]:
        """Execute tool by name with args. Returns (result, cost)."""
        # Track research activity
        if name in ("brave_search", "open", "click", "find", "interact"):
            self.research_actions += 1
        
        try:
            if name == "brave_search":
                return await self.brave_search(
                    query=args.get("query", ""),
                    response_length=args.get("response_length", "medium"),
                )
            
            elif name == "open":
                return await self.open(
                    ref_id_or_url=args.get("ref_id_or_url", ""),
                    start_line=args.get("start_line", 0),
                    response_length=args.get("response_length", "medium"),
                )
            
            elif name == "find":
                return await self.find(
                    ref_id=args.get("ref_id", ""),
                    pattern=args.get("pattern", ""),
                    response_length=args.get("response_length", "medium"),
                )
            
            elif name == "click":
                return await self.click(
                    ref_id=args.get("ref_id", ""),
                    link_id=args.get("link_id", 0),
                    response_length=args.get("response_length", "medium"),
                )
            
            elif name == "note":
                return self.note(content=args.get("content", ""))
            
            elif name == "list_files":
                return await self.list_files(directory=args.get("directory", "all"))
            
            elif name == "code_exec":
                return await self.code_exec(
                    script=args.get("script", ""),
                    description=args.get("description", ""),
                )
            
            elif name == "conclude_research":
                return self.conclude_research(summary=args.get("summary", ""))
            
            elif name == "breakdown":
                return self.breakdown(children=args.get("children", []))
            
            elif name == "submit_seed":
                return self.submit_seed(
                    ref_id=args.get("ref_id"),
                    lines=args.get("lines"),
                    content=args.get("content"),
                )
            
            elif name == "done":
                return self.done(reason=args.get("reason", ""))
            
            elif name == "interact":
                return await self.interact(
                    url_or_ref_id=args.get("url_or_ref_id", ""),
                    task=args.get("task", ""),
                )
            
            else:
                return f"Unknown tool: {name}", 0.0
                
        except Exception as e:
            logger.error(f"[ResearchTools] {name} failed: {e}")
            return f"Tool error: {e}", 0.0